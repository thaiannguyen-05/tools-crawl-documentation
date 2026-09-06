from __future__ import annotations

import io
import re
from pathlib import Path
from statistics import median

from markdown_it import MarkdownIt
from docx import Document
import pdfplumber

from .errors import UnsupportedFileError
from .file_fetcher import FileContent, fetch_file
from .types import ExtractedDocument, FileMetadata, Heading, Section

TEXT_EXTENSIONS = {".txt", ".text", ".md", ".markdown", ".pdf", ".docx"}


def extract_file(file_name: str, base_url: str | None = None) -> ExtractedDocument:
    file = fetch_file(file_name, base_url)
    return _extract(file_name, file)


def extract_text(file_name: str, content: bytes) -> ExtractedDocument:
    extension = Path(file_name).suffix.lower()
    file = FileContent(content=content, content_type="", size=len(content))
    return _extract(file_name, file)


def _extract(file_name: str, file: FileContent) -> ExtractedDocument:
    extension = Path(file_name).suffix.lower()

    if extension == ".pdf":
        sections, headings, page_count, para_count = _extract_pdf(file.content)
    elif extension == ".docx":
        sections, headings, page_count, para_count = _extract_docx(file.content)
    elif extension in (".md", ".markdown"):
        sections, headings, page_count, para_count = _extract_markdown(file.content)
    elif extension in (".txt", ".text"):
        sections, headings, page_count, para_count = _extract_txt(file.content)
    elif extension == ".doc":
        raise UnsupportedFileError(file_name, "legacy .doc is not supported, use .docx")
    else:
        raise UnsupportedFileError(
            file_name, f"supported extensions: {sorted(TEXT_EXTENSIONS)}"
        )

    metadata = FileMetadata(
        file_name=file_name,
        content_type=file.content_type,
        size=file.size,
        extension=extension,
        page_count=page_count,
        paragraph_count=para_count,
    )

    return ExtractedDocument(metadata=metadata, headings=headings, sections=sections)


ExtractionResult = tuple[list[Section], list[Heading], int | None, int]


def _extract_pdf(content: bytes) -> ExtractionResult:
    sections: list[Section] = []
    headings: list[Heading] = []
    page_count = 0
    para_count = 0

    try:
        pdf = pdfplumber.open(io.BytesIO(content))
    except Exception:
        return _extract_pdf_fallback(content)

    try:
        try:
            page_count = len(pdf.pages)
        except Exception:
            return _extract_pdf_fallback(content)

        all_sizes: list[float] = []
        for page in pdf.pages:
            for char in page.chars:
                all_sizes.append(char["size"])

        if not all_sizes:
            return [], [], page_count, 0

        body_size = median(all_sizes)
        heading_threshold = body_size * 1.5

        current_heading: Heading | None = None
        current_parts: list[str] = []

        for page in pdf.pages:
            words = page.extract_words(extra_attrs=["size"])
            line_groups: dict[float, list[str]] = {}
            line_sizes: dict[float, float] = {}

            for w in words:
                top = round(w["top"], 1)
                if top not in line_groups:
                    line_groups[top] = []
                    line_sizes[top] = 0.0
                line_groups[top].append(w["text"])
                line_sizes[top] = max(line_sizes[top], w["size"])

            for top in sorted(line_groups):
                line_text = " ".join(line_groups[top]).strip()
                line_size = line_sizes[top]

                if not line_text:
                    continue

                if line_size >= heading_threshold:
                    if current_parts:
                        text = "\n".join(current_parts).strip()
                        if text:
                            sections.append(
                                Section(heading=current_heading, text=text)
                            )
                        current_parts = []

                    level = _font_size_to_level(line_size, body_size)
                    heading = Heading(level=level, text=line_text)
                    headings.append(heading)
                    current_heading = heading
                else:
                    current_parts.append(line_text)
                    para_count += 1

        if current_parts:
            text = "\n".join(current_parts).strip()
            if text:
                sections.append(Section(heading=current_heading, text=text))
    except Exception:
        return _extract_pdf_fallback(content)
    finally:
        try:
            pdf.close()
        except Exception:
            pass

    return sections, headings, page_count, para_count


def _extract_pdf_fallback(content: bytes) -> ExtractionResult:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    sections: list[Section] = []
    para_count = 0

    for page in reader.pages:
        text = (page.extract_text() or "").strip()
        if text:
            sections.append(Section(heading=None, text=text))
            para_count += 1

    return sections, [], len(reader.pages), para_count


def _font_size_to_level(size: float, body_size: float) -> int:
    ratio = size / body_size
    if ratio >= 2.5:
        return 1
    if ratio >= 2.0:
        return 2
    return 3


def _extract_docx(content: bytes) -> ExtractionResult:
    document = Document(io.BytesIO(content))
    sections: list[Section] = []
    headings: list[Heading] = []
    current_heading: Heading | None = None
    current_parts: list[str] = []

    for paragraph in document.paragraphs:
        style_name = (paragraph.style.name or "").lower()

        if style_name.startswith("heading"):
            if current_parts:
                text = "\t".join(current_parts).strip()
                if text:
                    sections.append(Section(heading=current_heading, text=text))
                current_parts = []

            level = _parse_heading_level(style_name)
            heading = Heading(level=level, text=paragraph.text.strip())
            headings.append(heading)
            current_heading = heading
        else:
            current_parts.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            current_parts.append("\t".join(cell.text for cell in row.cells))

    if current_parts:
        text = "\t".join(current_parts).strip()
        if text:
            sections.append(Section(heading=current_heading, text=text))

    return sections, headings, None, len(document.paragraphs)


def _parse_heading_level(style_name: str) -> int:
    for part in style_name.split():
        if part.isdigit():
            return int(part)
    return 1


def _extract_markdown(content: bytes) -> ExtractionResult:
    source = content.decode("utf-8", errors="replace").strip()
    tokens = MarkdownIt().parse(source)

    sections: list[Section] = []
    headings: list[Heading] = []
    current_heading: Heading | None = None
    current_parts: list[str] = []
    para_count = 0

    for token in tokens:
        if token.type == "heading_open":
            if current_parts:
                text = "".join(current_parts).strip()
                if text:
                    sections.append(Section(heading=current_heading, text=text))
                current_parts = []

            level = int(token.tag[1]) if token.tag else 1

        elif token.type == "heading_close":
            heading_text = "".join(current_parts).strip() if current_parts else ""
            if token.tag:
                level = int(token.tag[1])
            else:
                level = 1
            heading = Heading(level=level, text=heading_text)
            headings.append(heading)
            current_heading = heading
            current_parts = []

        elif token.type == "inline" and token.children:
            for child in token.children:
                if child.type == "text":
                    current_parts.append(child.content)
        elif token.type in ("text", "code_block"):
            current_parts.append(token.content)
        elif token.type == "fence":
            current_parts.append(f"{token.content}\n")
        elif token.type == "hardbreak":
            current_parts.append("\n")
        elif token.type == "paragraph_close":
            para_count += 1
            text = "".join(current_parts).strip()
            if text:
                sections.append(Section(heading=current_heading, text=text))
                current_heading = None
            current_parts = []
        elif token.type in ("heading_close", "blockquote_close", "list_item_close"):
            pass

    if current_parts:
        text = "".join(current_parts).strip()
        if text:
            sections.append(Section(heading=current_heading, text=text))

    return sections, headings, None, para_count


_HEADING_UNDERLINE = re.compile(r"^[\-=]{3,}$")
_HEADING_CAPS = re.compile(r"^[A-Z][A-Z\s\d:.,\-]{2,60}$")
_NUM_RE = re.compile(r"^\d+(?:\.\d+)*$")


def _extract_txt(content: bytes) -> ExtractionResult:
    text = content.decode("utf-8", errors="replace").strip()
    if not text:
        return [], [], None, 0

    lines = text.split("\n")
    sections: list[Section] = []
    headings: list[Heading] = []
    current_heading: Heading | None = None
    current_parts: list[str] = []
    para_count = 0

    for line in lines:
        stripped = line.strip()

        if not stripped:
            if current_parts:
                para_count += 1
            continue

        heading_match = _detect_txt_heading(stripped)
        if heading_match:
            if current_parts:
                t = "\n".join(current_parts).strip()
                if t:
                    sections.append(Section(heading=current_heading, text=t))
                current_parts = []

            heading = Heading(level=heading_match[0], text=heading_match[1])
            headings.append(heading)
            current_heading = heading
        else:
            current_parts.append(stripped)

    if current_parts:
        t = "\n".join(current_parts).strip()
        if t:
            sections.append(Section(heading=current_heading, text=t))
            para_count += 1

    return sections, headings, None, para_count


def _detect_txt_heading(line: str) -> tuple[int, str] | None:
    parts = line.split(" ", 1)
    if len(parts) == 2:
        num_part, title = parts
        num_clean = num_part.rstrip(".)")
        if num_clean and _NUM_RE.match(num_clean) and len(title) >= 2:
            depth = num_clean.count(".") + 1
            return min(depth, 3), title

    if _HEADING_UNDERLINE.match(line):
        return None

    if _HEADING_CAPS.match(line):
        return 1, line

    return None
